"""Page générique d'une collection (CDthèque, Bibliothèque), pilotée par une configuration."""
import io
import json
import re
import uuid
from datetime import datetime

import pandas as pd
import requests
import streamlit as st

from common import (AVAILABLE, BUCKET, IMG_TYPES, UA, ask_ai, clean_url, db, ecoutes, flatten, fmt_eur, euro_number,
                    img_data_url, jpg_bytes, kind_of, parse_json, play, price_bounds, signed_url, to_excel,
                    year_number)

AUDIO_TYPES = ["mp3", "wav", "m4a", "ogg", "flac", "aac"]


class _Bytes:
    """Imite un fichier téléversé à partir d'octets (pour stocker une pochette trouvée en ligne)."""
    def __init__(self, data):
        self._d = data

    def getvalue(self):
        return self._d


def run(cfg):
    P = cfg["prefix"]
    F = cfg["fields"]                      # [(clé, libellé, type, options)]
    KEYS = [f[0] for f in F]
    LAB = {f[0]: f[1] for f in F} | {"created_at": "Date d'ajout", "Écoutes": "Écoutes"}
    TYPES = {f[0]: f[2] for f in F}
    OPTS = {f[0]: f[3] for f in F if len(f) > 3}
    PHOTOS = cfg["photos"]                 # [(colonne, nom de fichier, libellé)]
    CODE = cfg.get("code_field")
    sb = db()

    @st.cache_data(ttl=300, show_spinner=False)
    def load(table):
        return db().table(table).select("*").order("created_at", desc=True).execute().data or []

    def refresh():
        load.clear()
        st.rerun()

    try:
        records = load(cfg["table"])
    except Exception as e:
        st.error(f"Table « {cfg['table']} » inaccessible. Exécutez supabase_schema.sql dans Supabase.")
        st.code(str(e))
        st.stop()

    def prepare(recs):
        df = pd.DataFrame(recs)
        for c in KEYS + ["created_at", "id", "ecoutes"] + [p[0] for p in PHOTOS]:
            if c not in df.columns:
                df[c] = ""
        df = df.fillna("")
        df["_basse"] = df["estimation_basse_eur"].apply(euro_number)
        df["_haute"] = df["estimation_haute_eur"].apply(euro_number)
        df["_basse"], df["_haute"] = df["_basse"].fillna(df["_haute"]), df["_haute"].fillna(df["_basse"])
        df["_annee"] = df["annee"].apply(year_number)
        df["Écoutes"] = df["ecoutes"].apply(lambda v: len(ecoutes(v)))
        return df

    def label(r):
        main = r.get(cfg["main"]) or "Sans titre"
        return f"{main} — {r.get(cfg['sub'])}" if r.get(cfg["sub"]) else main

    def upload_photo(path, upload):
        sb.storage.from_(BUCKET).upload(path, jpg_bytes(upload), {"content-type": "image/jpeg", "upsert": "true"})

    def duplicates(vals, exclude=None):
        code = re.sub(r"\D", "", str(vals.get(CODE) or "")) if CODE else ""
        out = []
        for r in records:
            if r.get("id") == exclude:
                continue
            if code and re.sub(r"\D", "", str(r.get(CODE) or "")) == code:
                out.append(r)
            elif (vals.get(cfg["main"]) and str(r.get(cfg["main"]) or "").casefold() == vals[cfg["main"]].casefold()
                  and str(r.get(cfg["sub"]) or "").casefold() == str(vals.get(cfg["sub"]) or "").casefold()):
                out.append(r)
        return out

    # ------------------------------------------------------------ IA et recherche en ligne
    def normalize(data):
        flat = flatten(data)
        out = {k: flat.get(k, "") for k in KEYS}
        for k, alts in cfg.get("aliases", {}).items():
            out[k] = out[k] or next((flat[a] for a in alts if flat.get(a)), "")
        out["estimation_basse_eur"], out["estimation_haute_eur"] = price_bounds(out["estimation_basse_eur"],
                                                                                out["estimation_haute_eur"])
        for k, opts in OPTS.items():  # ramène les listes à une valeur proposée
            if out.get(k) and out[k] not in opts:
                out[k] = next((o for o in opts if o and o.casefold() in out[k].casefold()), out[k] if "" in opts else "")
        return out

    def identify(provider, uploads):
        ai_keys = [k for k in KEYS if k not in cfg.get("no_ai", [])]
        template = json.dumps({k: "" for k in ai_keys}, ensure_ascii=False)
        content = [{"type": "text", "text": cfg["prompt"] + "\nRéponds UNIQUEMENT avec ce JSON plat, valeurs en "
                                                           "chaînes de caractères :\n" + template}]
        for i, (u, (_, _, lab)) in enumerate(zip(uploads, PHOTOS)):
            if u:
                content += [{"type": "text", "text": f"PHOTO {i + 1} — {lab}"},
                            {"type": "image_url", "image_url": {"url": img_data_url(u)}}]
        raw = ask_ai(provider, content)
        st.session_state[f"{P}_raw"] = raw
        return normalize(parse_json(raw))

    def present(provider, rec):
        fiche = "\n".join(f"- {LAB[k]} : {rec.get(k)}" for k in KEYS
                          if rec.get(k) and k not in (cfg["text_field"], "notes", "identification_ia", "pistes"))
        if rec.get("pistes"):
            fiche += "\n- Pistes : " + str(rec["pistes"])[:1500]
        raw = ask_ai(provider, cfg["present_prompt"].format(fiche=fiche))
        try:
            return str(flatten(parse_json(raw)).get("texte", "")).strip()
        except Exception:
            return raw.strip()

    @st.cache_data(ttl=86400, show_spinner=False)
    def _lookup(table, code):  # la table fait partie de la clé de cache : CD et livres ne se mélangent pas
        return cfg["lookup"](code) if cfg.get("lookup") else {}

    def lookup(code):
        return _lookup(cfg["table"], code)

    def complete(provider, rec):
        """Complète une fiche sans photo (résumé/présentation en français, genre, estimation) à partir des infos connues."""
        keys = cfg.get("complete_fields", [cfg["text_field"], "estimation_basse_eur", "estimation_haute_eur"])
        fiche = "\n".join(f"- {LAB[k]} : {rec.get(k)}" for k in KEYS
                          if rec.get(k) and k not in (cfg["text_field"], "notes", "identification_ia", "pistes"))
        if rec.get("pistes"):
            fiche += "\n- Pistes : " + str(rec["pistes"])[:1500]
        template = json.dumps({k: "" for k in keys}, ensure_ascii=False)
        raw = ask_ai(provider, cfg["complete_prompt"].format(fiche=fiche) +
                     "\nRéponds UNIQUEMENT avec ce JSON, valeurs en chaînes de caractères :\n" + template)
        data = normalize(parse_json(raw))
        return {k: data.get(k, "") for k in keys}

    def apply_complete(provider, rec):
        """Ajoute à la fiche ce que l'IA propose : champs vides, plus ceux qu'elle sait mieux remplir (ex. genre en français)."""
        try:
            with st.spinner(f"Rédaction {cfg['text_art_de']} et estimation avec {provider}…"):
                extra = complete(provider, rec)
        except Exception as e:
            st.session_state[f"{P}_warn"] = f"Fiche remplie, mais l'IA n'a pas pu compléter : {e}"
            return rec
        out = dict(rec)
        for k, v in extra.items():
            if v and (not out.get(k) or k in cfg.get("complete_wins", [])):
                out[k] = v
        return out

    def merge(base, found):
        """Complète la fiche avec les données trouvées en ligne (prioritaires pour les faits, l'IA garde le reste)."""
        out = dict(base)
        for k, v in (found or {}).items():
            if k.startswith("_") or not v:
                continue
            if k in cfg.get("lookup_wins", []) or not out.get(k):
                out[k] = v
        return out

    # ------------------------------------------------------------ Formulaire
    def field_inputs(source, prefix):
        vals = {}
        cols = st.columns(2)
        shorts = [f for f in F if f[2] != "long"]
        longs = [f for f in F if f[2] == "long"]
        half = (len(shorts) + 1) // 2
        for i, f in enumerate(shorts):
            k = f[0]
            with cols[0 if i < half else 1]:
                cur = str(source.get(k) or "")
                if f[2] == "select":
                    opts = list(OPTS[k]) + ([cur] if cur and cur not in OPTS[k] else [])
                    vals[k] = st.selectbox(f[1], opts, index=opts.index(cur) if cur in opts else 0, key=f"{prefix}_{k}")
                else:
                    vals[k] = st.text_input(f[1], value=cur, key=f"{prefix}_{k}")
        for f in longs:
            k = f[0]
            vals[k] = st.text_area(f[1], value=str(source.get(k) or ""), height=140 if k in ("pistes", cfg["text_field"]) else 80,
                                   key=f"{prefix}_{k}")
        return vals

    # ------------------------------------------------------------ Écoutes (CD)
    def save_ecoutes(rid, items, msg):
        sb.table(cfg["table"]).update({"ecoutes": json.dumps(items, ensure_ascii=False) if items else None}) \
            .eq("id", rid).execute()
        st.toast(msg)
        refresh()

    def ecoute_panel(rec):
        rid = rec["id"]
        items = ecoutes(rec.get("ecoutes"))
        st.markdown(f"#### 🎧 Écouter ({len(items)})" if items else "#### 🎧 Écouter")
        for i, item in enumerate(items):
            a, b = st.columns([6, 1])
            a.caption(f"{i + 1}. {item.get('nom') or kind_of(item['url'])}")
            if b.button("🗑️", key=f"{P}_rm_{rid}_{i}", help="Supprimer cette écoute"):
                if not item["url"].startswith("http"):
                    try:
                        sb.storage.from_(BUCKET).remove([item["url"]])
                    except Exception:
                        pass
                save_ecoutes(rid, items[:i] + items[i + 1:], "Écoute supprimée.")
            play(item)
        q = requests.utils.quote(f"{rec.get('artiste') or ''} {rec.get('titre') or ''}".strip())
        if q:
            st.caption(f"Chercher : [YouTube](https://www.youtube.com/results?search_query={q}) · "
                       f"[Spotify](https://open.spotify.com/search/{q}) · [Deezer](https://www.deezer.com/search/{q})")
        with st.expander("➕ Ajouter une écoute"):
            with st.form(f"{P}_lnk_{rid}", clear_on_submit=True, border=False):
                link = st.text_input("Lien (YouTube, Spotify, Deezer, MP3…)")
                nom = st.text_input("Nom (facultatif)", placeholder="ex. album complet, piste 3…")
                if st.form_submit_button("💾 Enregistrer ce lien", type="primary"):
                    url = clean_url(link)
                    if not url.startswith("http"):
                        st.error("Le lien doit commencer par http.")
                    elif any(clean_url(i["url"]) == url for i in items):
                        st.warning("Ce lien est déjà enregistré.")
                    else:
                        save_ecoutes(rid, items + [{"url": url, "nom": nom.strip() or kind_of(url),
                                                    "ajout": datetime.now().strftime("%Y-%m-%d")}], "Écoute ajoutée.")
            up = st.file_uploader("Ou déposez un fichier audio", type=AUDIO_TYPES, key=f"{P}_aud_{rid}")
            if up and st.button("💾 Enregistrer ce fichier", key=f"{P}_audb_{rid}"):
                ext = up.name.rsplit(".", 1)[-1].lower()
                path = f"{P}/{rid}/audio_{uuid.uuid4().hex[:6]}.{ext}"
                sb.storage.from_(BUCKET).upload(path, up.getvalue(), {"content-type": up.type or f"audio/{ext}",
                                                                      "upsert": "true"})
                save_ecoutes(rid, items + [{"url": path, "nom": up.name.rsplit(".", 1)[0]}], "Fichier ajouté.")

    # ------------------------------------------------------------ En-tête
    h1, h2 = st.columns([6, 1])
    with h1:
        st.title(cfg["title"])
        st.caption(cfg["caption"])
    with h2:
        st.write("")
        if st.button("🔄 Actualiser", key=f"{P}_refresh"):
            refresh()
    tabs = st.tabs([f"➕ Ajouter {cfg['un']}", "📚 Catalogue", "📊 Récapitulatif", "⬇️ Exporter"])

    # ------------------------------------------------------------ Ajout
    with tabs[0]:
        n = st.session_state.setdefault(f"{P}_nonce", 0)
        if st.session_state.pop(f"{P}_saved", False):
            st.success(f"{cfg['item'][0].upper() + cfg['item'][1:]} enregistré. Vous pouvez passer au suivant.")
        st.write(cfg["add_help"])
        pcols = st.columns(len(PHOTOS))
        uploads = []
        for (col, fname, lab), pc in zip(PHOTOS, pcols):
            with pc:
                u = st.file_uploader(f"📷 {lab}", type=IMG_TYPES, key=f"{P}_up_{fname}_{n}")
                if u:
                    st.image(u, width=220)
                uploads.append(u)

        result = st.session_state.get(f"{P}_result", {})
        provider = None
        if len(AVAILABLE) > 1:
            provider = st.radio("IA utilisée", list(AVAILABLE), horizontal=True, key=f"{P}_provider")
        elif AVAILABLE:
            provider = next(iter(AVAILABLE))
        if CODE:
            c1, c2 = st.columns([3, 1])
            with c1:
                code_in = st.text_input(f"🔢 {LAB[CODE]} (facultatif)", key=f"{P}_code_{n}",
                                        placeholder=cfg.get("code_placeholder", ""))
            with c2:
                st.write("")
                st.write("")
                do_lookup = st.button(f"🔎 Rechercher en ligne", key=f"{P}_lk_{n}",
                                      disabled=not re.sub(r"\D", "", code_in or ""))
            if do_lookup:
                code = re.sub(r"[^\dXx]", "", code_in)
                try:
                    with st.spinner(f"Recherche sur {cfg['lookup_name']}…"):
                        found = lookup(code)
                    if found:
                        filled = merge({**result, CODE: code}, found)
                        if provider and not filled.get(cfg["text_field"]) or provider and cfg.get("complete_wins"):
                            filled = apply_complete(provider, filled)
                        st.session_state[f"{P}_result"] = filled
                        st.session_state[f"{P}_cover"] = found.get("_cover")
                        st.session_state[f"{P}_round"] = st.session_state.get(f"{P}_round", 0) + 1
                        st.rerun()
                    else:
                        st.warning(f"Rien trouvé sur {cfg['lookup_name']} pour {code}. Utilisez l'analyse photo "
                                   "ou saisissez la fiche à la main.")
                except Exception as e:
                    st.error(f"Recherche impossible : {e}")

        has_photo, has_info = any(uploads), bool((result.get(cfg["main"]) or "").strip())
        ai_label = ("🤖 Analyser les photos avec l'IA" if has_photo else
                    f"🤖 Compléter la fiche avec l'IA ({cfg['text_art']}, estimation)")
        if st.session_state.get(f"{P}_warn"):
            st.warning(st.session_state.pop(f"{P}_warn"))
        if provider and not has_photo and not has_info:
            st.caption("Ajoutez une photo, ou faites d'abord la recherche en ligne, pour activer l'IA.")
        if st.button(ai_label, type="primary", key=f"{P}_ai_{n}",
                     disabled=not (provider and (has_photo or has_info))):
            if not has_photo:
                st.session_state[f"{P}_result"] = apply_complete(provider, result)
                st.session_state[f"{P}_round"] = st.session_state.get(f"{P}_round", 0) + 1
                st.rerun()
            try:
                with st.spinner(f"Analyse avec {provider}…"):
                    res = identify(provider, uploads)
                    code = re.sub(r"[^\dXx]", "", res.get(CODE) or "") if CODE else ""
                    if code and len(code) >= 10:
                        try:
                            found = lookup(code)
                            res = merge(res, found)
                            st.session_state[f"{P}_cover"] = found.get("_cover") if found else None
                        except Exception:
                            pass
                st.session_state[f"{P}_result"] = {**result, **{k: v for k, v in res.items() if v}}
                st.session_state[f"{P}_round"] = st.session_state.get(f"{P}_round", 0) + 1
                st.rerun()
            except Exception as e:
                st.error(f"Erreur IA : {e}")
        if st.session_state.get(f"{P}_raw"):
            with st.expander("Réponse brute de l'IA (diagnostic)"):
                st.code(st.session_state[f"{P}_raw"])

        if any(result.values()):
            st.success("Fiche pré-remplie : relisez et corrigez avant d'enregistrer.")
            dups = duplicates(result)
            if dups:
                st.warning(f"⚠️ Doublon possible : « {label(dups[0])} » est déjà dans la collection.")
            if st.session_state.get(f"{P}_cover") and not uploads[0]:
                st.image(st.session_state[f"{P}_cover"], width=160, caption="Couverture trouvée en ligne "
                                                                             "(sera enregistrée si vous n'ajoutez pas de photo)")

        with st.form(f"{P}_new_{n}"):
            st.subheader(f"Fiche {cfg['du']}")
            vals = field_inputs(result, f"{P}_new{n}_{st.session_state.get(f'{P}_round', 0)}")
            submitted = st.form_submit_button(f"💾 Enregistrer {cfg['le']}", type="primary")
        if submitted:
            if not (vals.get(cfg["main"]) or "").strip():
                st.error(f"Renseignez au moins : {LAB[cfg['main']]}.")
            else:
                try:
                    with st.spinner("Enregistrement…"):
                        rid = str(uuid.uuid4())
                        row = {"id": rid, "created_at": datetime.now().astimezone().isoformat(),
                               **{k: (v.strip() if isinstance(v, str) else v) for k, v in vals.items()}}
                        for (col, fname, _), u in zip(PHOTOS, uploads):
                            if u:
                                row[col] = f"{P}/{rid}/{fname}.jpg"
                                upload_photo(row[col], u)
                        cover = st.session_state.get(f"{P}_cover")
                        if cover and not uploads[0]:
                            try:
                                r = requests.get(cover, headers=UA, timeout=20)
                                if r.ok and len(r.content) > 2000:
                                    path = f"{P}/{rid}/{PHOTOS[0][1]}.jpg"
                                    upload_photo(path, _Bytes(r.content))
                                    row[PHOTOS[0][0]] = path  # seulement si l'image a bien été stockée
                            except Exception:
                                pass
                        sb.table(cfg["table"]).insert(row).execute()
                    for k in ("result", "cover", "raw"):
                        st.session_state.pop(f"{P}_{k}", None)
                    st.session_state[f"{P}_nonce"] = n + 1
                    st.session_state[f"{P}_saved"] = True
                    refresh()
                except Exception as e:
                    st.error(f"Enregistrement impossible : {e}")
        if any(result.values()) and st.button("Effacer la fiche en cours", key=f"{P}_clear_{n}"):
            for k in ("result", "cover", "raw"):
                st.session_state.pop(f"{P}_{k}", None)
            st.session_state[f"{P}_nonce"] = n + 1
            st.rerun()

    # ------------------------------------------------------------ Catalogue
    with tabs[1]:
        if not records:
            st.info(f"Aucun {cfg['item']} enregistré pour l'instant.")
        else:
            df = prepare(records)
            fc = st.columns([3] + [2] * len(cfg["filters"]))
            with fc[0]:
                q = st.text_input("🔍 Rechercher", key=f"{P}_q", placeholder=cfg["search_hint"])
            sel = {}
            for k, c in zip(cfg["filters"], fc[1:]):
                with c:
                    values = sorted({str(x) for x in df[k] if str(x).strip()}, key=str.casefold)
                    sel[k] = st.multiselect(LAB[k], values, key=f"{P}_f_{k}")
            if q:
                mask = df[KEYS].astype(str).apply(lambda c: c.str.contains(q, case=False, regex=False, na=False))
                df = df[mask.any(axis=1)]
            for k, v in sel.items():
                if v:
                    df = df[df[k].astype(str).isin(v)]
            o1, o2 = st.columns([3, 2])
            sort_keys = ["created_at"] + cfg["sort_keys"]
            with o1:
                sort_by = st.selectbox("Trier par", sort_keys, format_func=LAB.get, key=f"{P}_sort")
            with o2:
                desc = st.toggle("Ordre décroissant", value=(sort_by == "created_at"), key=f"{P}_desc_{sort_by}")
            ser = {"annee": df["_annee"], "estimation_basse_eur": df["_basse"], "estimation_haute_eur": df["_haute"],
                   "created_at": pd.to_datetime(df["created_at"], errors="coerce", utc=True)}.get(
                sort_by, df[sort_by].astype(str).str.casefold().where(lambda x: x != ""))
            df = df.assign(_s=ser).sort_values("_s", ascending=not desc, na_position="last", kind="stable")
            st.caption(f"{len(df)} {cfg['item_pl']} affiché(s) sur {len(records)} · valeur estimée "
                       f"{fmt_eur(df['_basse'].sum(min_count=1))} – {fmt_eur(df['_haute'].sum(min_count=1))}")
            cols = cfg["table_cols"]
            ccfg = {LAB[k]: st.column_config.TextColumn(LAB[k], width="large") for k in cols if TYPES.get(k) == "long"}
            event = st.dataframe(df[cols].rename(columns=LAB), width="stretch", hide_index=True, column_config=ccfg,
                                 on_select="rerun", selection_mode="single-row", key=f"{P}_table")
            st.caption("Cliquez sur une ligne pour ouvrir la fiche.")
            rows = event.selection.rows if event and hasattr(event, "selection") else []
            if rows and rows[0] < len(df):
                rec = df.iloc[rows[0]].to_dict()
                rid = rec["id"]
                st.divider()
                st.subheader(f"{cfg['icon']} {rec.get(cfg['main']) or 'Sans titre'}")
                if rec.get(cfg["sub"]):
                    st.markdown(f"**{rec[cfg['sub']]}**")
                pics, info = st.columns([1, 2])
                with pics:
                    for col, _, lab in PHOTOS:
                        if rec.get(col):
                            try:
                                st.image(signed_url(rec[col]), caption=lab, width="stretch")
                            except Exception:
                                st.caption(f"{lab} : photo indisponible")
                with info:
                    for k in KEYS:
                        if k in (cfg["main"], cfg["sub"]) or TYPES[k] == "long" or not rec.get(k):
                            continue
                        st.markdown(f"**{LAB[k]} :** {rec[k]}")
                    tf = cfg["text_field"]
                    if rec.get(tf):
                        st.info(rec[tf], icon="📝")
                    if AVAILABLE and st.button(f"✨ {'Refaire' if rec.get(tf) else 'Rédiger'} {cfg['text_art']} "
                                               f"avec l'IA", key=f"{P}_present_{rid}"):
                        try:
                            prov = st.session_state.get(f"{P}_provider") or next(iter(AVAILABLE))
                            with st.spinner(f"Rédaction avec {prov}…"):
                                txt = present(prov, rec)
                            if txt:
                                sb.table(cfg["table"]).update({tf: txt}).eq("id", rid).execute()
                                refresh()
                            st.warning("L'IA n'a rien pu écrire.")
                        except Exception as e:
                            st.error(f"Rédaction impossible : {e}")
                    for k in KEYS:
                        if TYPES[k] == "long" and k != tf and rec.get(k):
                            with st.expander(LAB[k], expanded=(k == "pistes")):
                                st.text(rec[k])
                if cfg.get("ecoutes"):
                    ecoute_panel(rec)

                with st.expander("✏️ Modifier cette fiche"):
                    with st.form(f"{P}_edit_{rid}"):
                        ev = field_inputs(rec, f"{P}_e_{rid}")
                        st.markdown("**Remplacer une photo (facultatif)**")
                        rc = st.columns(len(PHOTOS))
                        newp = {col: rc[i].file_uploader(lab, type=IMG_TYPES, key=f"{P}_rp_{col}_{rid}")
                                for i, (col, _, lab) in enumerate(PHOTOS)}
                        ok = st.form_submit_button("Enregistrer les modifications", type="primary")
                    if ok:
                        try:
                            upd = {k: (v.strip() if isinstance(v, str) else v) for k, v in ev.items()}
                            for col, fname, _ in PHOTOS:
                                if newp[col]:
                                    path = rec.get(col) or f"{P}/{rid}/{fname}.jpg"
                                    upload_photo(path, newp[col])
                                    upd[col] = path
                                    signed_url.clear()
                            sb.table(cfg["table"]).update(upd).eq("id", rid).execute()
                            st.toast("Modifications enregistrées.")
                            refresh()
                        except Exception as e:
                            st.error(f"Modification impossible : {e}")
                with st.expander("🗑️ Supprimer cette fiche"):
                    sure = st.checkbox(f"Je confirme vouloir supprimer {cfg['ce']}", key=f"{P}_sure_{rid}")
                    if st.button("Supprimer définitivement", disabled=not sure, key=f"{P}_del_{rid}"):
                        try:
                            paths = [rec[c] for c, _, _ in PHOTOS if rec.get(c)]
                            paths += [i["url"] for i in ecoutes(rec.get("ecoutes")) if not i["url"].startswith("http")]
                            if paths:
                                try:
                                    sb.storage.from_(BUCKET).remove(paths)
                                except Exception:
                                    pass
                            sb.table(cfg["table"]).delete().eq("id", rid).execute()
                            st.toast("Fiche supprimée.")
                            refresh()
                        except Exception as e:
                            st.error(f"Suppression impossible : {e}")

    # ------------------------------------------------------------ Récapitulatif
    with tabs[2]:
        if not records:
            st.info("Le récapitulatif apparaîtra dès le premier ajout.")
        else:
            rc = prepare(records)
            mets = [(cfg["item_pl"][0].upper() + cfg["item_pl"][1:], len(rc)),
                    ("Valeur estimée basse", fmt_eur(rc["_basse"].sum(min_count=1))),
                    ("Valeur estimée haute", fmt_eur(rc["_haute"].sum(min_count=1)))] + \
                [(lab, fn(rc)) for lab, fn in cfg.get("extra_metrics", [])]
            for (lab, v), c in zip(mets, st.columns(len(mets))):
                c.metric(lab, v)
            st.caption("Montants indicatifs : ni expertise ni garantie de prix.")
            groups = cfg["recap_by"]
            for i in range(0, len(groups), 2):
                cs = st.columns(2)
                for k, c in zip(groups[i:i + 2], cs):
                    with c:
                        st.markdown(f"#### Par {LAB[k].lower()}")
                        s = rc[k].replace("", "Non renseigné").value_counts().head(15)
                        st.bar_chart(s.rename_axis(LAB[k]).reset_index(name="Nombre"), x=LAB[k], y="Nombre",
                                     horizontal=True)
            dec = rc["_annee"].dropna().astype(int) // 10 * 10
            if len(dec):
                st.markdown("#### Par décennie")
                st.bar_chart((dec.astype(str) + "s").value_counts().sort_index().rename_axis("Décennie")
                             .reset_index(name="Nombre"), x="Décennie", y="Nombre")

    # ------------------------------------------------------------ Export
    def fetch_photo(path):
        """Octets d'une photo stockée : téléchargement direct, sinon via un lien signé."""
        if not path:
            return None
        try:
            data = db().storage.from_(BUCKET).download(path)
            if isinstance(data, (bytes, bytearray)) and len(data) > 100:
                return bytes(data)
        except Exception:
            pass
        try:
            r = requests.get(signed_url(path), timeout=30)
            if r.ok and len(r.content) > 100:
                return r.content
        except Exception:
            pass
        return None

    @st.cache_data(ttl=3600, show_spinner=False, max_entries=1000)
    def photo_bytes(path):
        return fetch_photo(path)

    def excel_with_photos(df, photo_col):
        """Classeur Excel avec une miniature de la photo dans la première colonne de chaque ligne."""
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from PIL import Image as PILImage
        wb = Workbook()
        ws = wb.active
        ws.title = cfg["pdf_title"][:31]
        cols = [c for c in df.columns if not str(c).startswith("_") and c != photo_col]
        ws.append(["Photo"] + cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F1A17")
        ws.column_dimensions["A"].width = 14
        for i, col in enumerate(cols, start=2):
            letter = ws.cell(row=1, column=i).column_letter
            wide = col in (LAB.get(cfg["text_field"]), LAB.get("pistes"), LAB.get("notes"))
            ws.column_dimensions[letter].width = 60 if wide else 18
        for r_i, row in enumerate(df.to_dict("records"), start=2):
            ws.append([""] + [str(row.get(c) or "") for c in cols])
            ws.row_dimensions[r_i].height = 84
            for cell in ws[r_i]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            data = photo_bytes(row.get(photo_col)) if row.get(photo_col) else None
            if data:
                try:
                    im = PILImage.open(io.BytesIO(data)).convert("RGB")
                    im.thumbnail((140, 110))
                    b = io.BytesIO()
                    im.save(b, "PNG")
                    b.seek(0)
                    xl = XLImage(b)
                    ws.add_image(xl, f"A{r_i}")
                except Exception:
                    pass
        ws.freeze_panes = "B2"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    with tabs[3]:
        if not records:
            st.info(f"Ajoutez {cfg['un']} pour commencer.")
        else:
            st.markdown("### 📖 Édition PDF du catalogue")
            st.caption(f"Une fiche par {cfg['item']} avec la photo, les informations et {cfg['text_art']} — "
                       "à imprimer ou à envoyer.")
            k1, k2, k3 = st.columns(3)
            with k1:
                wph = st.toggle("Avec les photos", value=True, key=f"{P}_cph2")
            with k2:
                wt = st.toggle(f"Avec {cfg['text_art']}", value=True, key=f"{P}_ct2")
            with k3:
                wp = st.toggle("Avec les prix estimés", value=True, key=f"{P}_cp2")
            k4, k5 = st.columns(2)
            with k4:
                title = st.text_input("Titre du catalogue", value=cfg["pdf_title"], key=f"{P}_ctitle")
            with k5:
                order = st.selectbox("Ordre", cfg["sort_keys"], format_func=LAB.get, key=f"{P}_corder")
            cdf = prepare(records)
            cdf = cdf.assign(_s=cdf["_annee"] if order == "annee" else
                             cdf[order].astype(str).str.casefold().where(lambda x: x != "")) \
                     .sort_values("_s", na_position="last", kind="stable")
            photo_col, tf = cfg["pdf"]["photo"], cfg["text_field"]
            n_ph = int((cdf[photo_col].astype(str).str.strip() != "").sum())
            n_tx = int((cdf[tf].astype(str).str.strip() != "").sum())
            st.caption(f"{len(cdf)} fiche(s) · {n_ph} avec photo · {n_tx} avec {cfg['text_art'].split(' ', 1)[1]}.")
            if st.button("📄 Créer l'édition PDF", type="primary", key=f"{P}_cbuild"):
                try:
                    from catalogue_pdf import build_generic
                    got, missing = [], []

                    def loader(path):
                        data = photo_bytes(path)
                        (got if data else missing).append(path)
                        return data
                    with st.spinner("Mise en page" + (" et récupération des photos…" if wph else "…")):
                        st.session_state[f"{P}_pdf"] = build_generic(cdf, cfg["pdf"], wp, wph, wt, title,
                                                                     load_photo=loader)
                        st.session_state[f"{P}_pdf_opts"] = (wp, wph, wt, len(got), len(missing))
                except Exception as e:
                    st.error(f"Édition PDF impossible : {e}")
            if st.session_state.get(f"{P}_pdf"):
                pr, ph, tx, ok_n, miss_n = st.session_state[f"{P}_pdf_opts"]
                name = f"catalogue_{cfg['table']}" + ("" if pr else "_sans_prix") + ("" if ph else "_sans_photos")
                st.download_button("⬇️ Télécharger l'édition PDF", st.session_state[f"{P}_pdf"],
                                   f"{name}_{datetime.now():%Y-%m-%d}.pdf", "application/pdf", key=f"{P}_cdl",
                                   type="primary")
                st.caption(f"Contenu : {'avec' if ph else 'sans'} photos ({ok_n} incluses) · "
                           f"{'avec' if tx else 'sans'} {cfg['text_art'].split(' ', 1)[1]} · {'avec' if pr else 'sans'} prix.")
                if miss_n:
                    st.warning(f"{miss_n} photo(s) n'ont pas pu être récupérées depuis le stockage.")

            st.divider()
            st.markdown("### 📊 Tableur")
            prep = prepare(records)
            out = prep[KEYS + ["created_at"] + [photo_col]].rename(columns={k: v for k, v in LAB.items() if k != photo_col})
            if cfg.get("ecoutes"):
                out["Écoutes"] = [", ".join(i["url"] for i in ecoutes(v) if i["url"].startswith("http"))
                                  for v in prep["ecoutes"]]
            stamp = datetime.now().strftime("%Y-%m-%d")
            x1, x2, x3 = st.columns(3)
            with x1:
                if st.button("📗 Préparer l'Excel avec photos", key=f"{P}_xlp"):
                    with st.spinner("Insertion des photos dans le classeur…"):
                        st.session_state[f"{P}_xlsx"] = excel_with_photos(out, photo_col)
                if st.session_state.get(f"{P}_xlsx"):
                    st.download_button("⬇️ Excel avec photos", st.session_state[f"{P}_xlsx"],
                                       f"{cfg['table']}_photos_{stamp}.xlsx",
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       key=f"{P}_xlpd")
            plain = out.drop(columns=[photo_col])
            x2.download_button("📗 Excel simple (.xlsx)", to_excel(plain), f"{cfg['table']}_{stamp}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"{P}_xl")
            x3.download_button("📄 CSV", plain.to_csv(index=False, sep=";").encode("utf-8-sig"),
                               f"{cfg['table']}_{stamp}.csv", "text/csv", key=f"{P}_csv")
